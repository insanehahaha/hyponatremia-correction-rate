"""Word table output in journal style (no vertical rules, single top and bottom rules)."""
import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

DASH = "–"      # en dash for ranges
EMPTY = "—"     # em dash for not applicable


def is_empty(v):
    return v is None or (isinstance(v, float) and np.isnan(v)) or v == ""


def ci(v, a, b, nd=2):
    return f"{v:.{nd}f} ({a:.{nd}f}{DASH}{b:.{nd}f})"


def ci_to(v, a, b, nd=1):
    return f"{v:.{nd}f} ({a:.{nd}f} to {b:.{nd}f})"


def section_row(row):
    """A row whose cells after the label are all empty is rendered as a merged section heading."""
    return all(is_empty(v) for v in row.iloc[1:])


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9)
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = ""
    props.comments = ""
    props.revision = 1
    return doc


def _format_cell(cell, size, bold=False, italic=False):
    for p in cell.paragraphs:
        pf = p.paragraph_format
        pf.line_spacing = 1.0
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        for r in p.runs:
            r.font.size = Pt(size)
            r.bold = bold
            r.italic = italic


def _borders(table):
    """Single rule above the table, below the header row and below the table; no vertical rules."""
    tblPr = table._tbl.tblPr
    for e in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(e)
    borders = OxmlElement("w:tblBorders")
    for name, val in [("top", "single"), ("left", "nil"), ("bottom", "single"), ("right", "nil"), ("insideH", "nil"), ("insideV", "nil")]:
        e = OxmlElement(f"w:{name}")
        e.set(qn("w:val"), val)
        if val != "nil":
            e.set(qn("w:sz"), "8")
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)
    for c in table.rows[0].cells:
        tcPr = c._tc.get_or_add_tcPr()
        bd = OxmlElement("w:tcBorders")
        e = OxmlElement("w:bottom")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        bd.append(e)
        tcPr.append(bd)


def _cell_text(column, value):
    if is_empty(value):
        return ""
    if isinstance(value, float):
        return f"{value:.2f}" if str(column).startswith(("Largest SMD", "SMD")) else f"{value:.1f}"
    return str(value)


def add_table(doc, df, title, footnote=None, section=None, widths=None):
    """Append a titled table; rows for which `section(row)` is true become merged bold-italic headings."""
    doc.add_paragraph(title).runs[0].bold = True
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"
    t.autofit = widths is None
    for j, c in enumerate(df.columns):
        cell = t.rows[0].cells[j]
        cell.text = str(c)
        _format_cell(cell, 8, bold=True)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        if section is not None and section(row):
            merged = cells[0].merge(cells[-1])
            merged.text = str(row.iloc[0])
            _format_cell(merged, 8, bold=True, italic=True)
            continue
        for j, (c, v) in enumerate(zip(df.columns, row)):
            cells[j].text = _cell_text(c, v)
            _format_cell(cells[j], 8)
    _borders(t)
    if widths is not None:
        for r in t.rows:
            for j, w in enumerate(widths):
                try:
                    r.cells[j].width = Inches(w)
                except IndexError:
                    pass
    if footnote:
        p = doc.add_paragraph(footnote)
        p.runs[0].font.size = Pt(8)
        p.paragraph_format.line_spacing = 1.0
    doc.add_paragraph()
